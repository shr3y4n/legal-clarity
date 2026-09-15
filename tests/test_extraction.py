import io
import pytest
from fastapi import HTTPException
import pypdf
import docx
from app.services.extraction.extractor import process_document
from app.services.extraction.txt_extractor import extract_txt


def create_sample_pdf() -> bytes:
    from reportlab.pdfgen import canvas
    buf = io.BytesIO()
    c = canvas.Canvas(buf)
    c.drawString(72, 750, "SECTION 1.0 TERM AND TERMINATION")
    c.drawString(72, 730, "Either party may terminate upon 30 days prior written notice.")
    c.save()
    return buf.getvalue()


def create_sample_docx() -> bytes:
    doc = docx.Document()
    doc.add_heading("NON-DISCLOSURE AGREEMENT", level=1)
    p = doc.add_paragraph("ARTICLE 1. CONFIDENTIALITY OBLIGATIONS")
    p.add_run("\nThe recipient shall protect proprietary data with reasonable care.")
    doc.add_paragraph("SECTION 2. TERM\nThis agreement remains in effect for 3 years.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_txt_extraction_with_sections():
    sample_text = (
        "RESIDENTIAL LEASE AGREEMENT\n\n"
        "SECTION 1.0 PARTIES AND PREMISES\n"
        "Landlord agrees to lease to Tenant the premises at 123 Elm Street.\n\n"
        "SECTION 2.0 RENT PAYMENT\n"
        "Rent shall be $2,500 per month payable on the 1st calendar day of each month.\n\n"
        "SECTION 3.0 NOTICE PERIOD\n"
        "Tenant must provide 60 days written notice prior to departure.\n"
    ).encode("utf-8")

    doc = process_document("lease.txt", sample_text, "text/plain")
    assert doc.metadata.filename == "lease.txt"
    assert doc.metadata.page_count >= 1
    assert len(doc.pages) >= 1
    assert "RESIDENTIAL LEASE AGREEMENT" in doc.full_text
    assert len(doc.pages[0].sections) >= 2


def test_docx_extraction_valid():
    docx_bytes = create_sample_docx()
    doc = process_document("nda.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert doc.metadata.filename == "nda.docx"
    assert "CONFIDENTIALITY OBLIGATIONS" in doc.full_text
    assert doc.metadata.page_count >= 1


def test_pdf_extraction_valid():
    pdf_bytes = create_sample_pdf()
    doc = process_document("contract.pdf", pdf_bytes, "application/pdf")
    assert doc.metadata.filename == "contract.pdf"
    assert doc.metadata.page_count == 1
    assert "SECTION 1.0" in doc.full_text


def test_unicode_and_special_characters():
    unicode_text = (
        "SECTION 1. JURISDICTION & LIQUIDATED DAMAGES\n"
        "Damages: €50,000 or ¥5,000,000. Governing Law: Zürich, Switzerland. Special chars: £, ©, ®, § 4.2.\n"
    ).encode("utf-8")
    doc = process_document("intl_agreement.txt", unicode_text, "text/plain")
    assert "€50,000" in doc.full_text
    assert "Zürich" in doc.full_text
    assert "§ 4.2" in doc.full_text


def test_empty_document_rejection():
    with pytest.raises(HTTPException) as exc_info:
        process_document("empty.txt", b"", "text/plain")
    assert exc_info.value.status_code == 400


def test_corrupt_docx_rejection():
    fake_docx = b"PK\x03\x04corrupted_payload_without_xml"
    with pytest.raises(HTTPException) as exc_info:
        process_document("bad.docx", fake_docx, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")
    assert exc_info.value.status_code == 400
