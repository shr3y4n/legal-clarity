import io
import re
from typing import List, Tuple

import docx

from app.models.schemas import Page, Section

_SECTION_PATTERN = re.compile(
    r"^(?:(?:ARTICLE|SECTION|CLAUSE)\s+([0-9A-Z\.]+)|([0-9]+\.[0-9]+(?:\.[0-9]+)?))\s*[:\.\-]?\s*(.*)",
    re.IGNORECASE
)

# Approximate words per standard page in formatted legal contracts
WORDS_PER_PAGE = 350


def extract_docx(file_bytes: bytes) -> Tuple[List[Page], str]:
    """
    Extracts text and structure from DOCX bytes, mapping paragraphs and tables into
    deterministic pages and sections.
    """
    doc = docx.Document(io.BytesIO(file_bytes))
    paragraphs_text: List[Tuple[str, str, bool]] = []  # (text, style_name, is_heading)

    for p in doc.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        is_heading = p.style.name.startswith("Heading") or bool(_SECTION_PATTERN.match(t))
        paragraphs_text.append((t, p.style.name, is_heading))

    # Also extract table text
    for table in doc.tables:
        for row in table.rows:
            row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_cells:
                paragraphs_text.append((" | ".join(row_cells), "Table", False))

    if not paragraphs_text:
        return [], ""

    pages: List[Page] = []
    current_page_text_parts: List[str] = []
    current_sections: List[Section] = []
    current_page_num = 1
    current_word_count = 0

    current_heading = None
    current_clause = None
    current_sec_lines: List[str] = []

    for text, style_name, is_heading in paragraphs_text:
        words = len(text.split())
        match = _SECTION_PATTERN.match(text)

        if is_heading or match:
            # Flush previous section
            if current_sec_lines:
                sec_text = "\n".join(current_sec_lines).strip()
                if sec_text:
                    current_sections.append(
                        Section(
                            section_id=f"p{current_page_num}_s{len(current_sections)+1}",
                            page_number=current_page_num,
                            heading=current_heading or f"Section {len(current_sections)+1}",
                            clause_number=current_clause,
                            text=sec_text,
                            start_char=0,
                            end_char=len(sec_text)
                        )
                    )
            clause_part = match.group(1) or match.group(2) if match else None
            current_clause = clause_part.strip() if clause_part else None
            current_heading = text
            current_sec_lines = [text]
        else:
            current_sec_lines.append(text)

        current_page_text_parts.append(text)
        current_word_count += words

        # Page threshold
        if current_word_count >= WORDS_PER_PAGE:
            if current_sec_lines:
                sec_text = "\n".join(current_sec_lines).strip()
                if sec_text:
                    current_sections.append(
                        Section(
                            section_id=f"p{current_page_num}_s{len(current_sections)+1}",
                            page_number=current_page_num,
                            heading=current_heading or f"Page {current_page_num}",
                            clause_number=current_clause,
                            text=sec_text,
                            start_char=0,
                            end_char=len(sec_text)
                        )
                    )
                current_sec_lines = []
            page_full_text = "\n\n".join(current_page_text_parts)
            pages.append(Page(page_number=current_page_num, text=page_full_text, sections=current_sections))
            current_page_num += 1
            current_page_text_parts = []
            current_sections = []
            current_word_count = 0

    # Flush remaining
    if current_sec_lines:
        sec_text = "\n".join(current_sec_lines).strip()
        if sec_text:
            current_sections.append(
                Section(
                    section_id=f"p{current_page_num}_s{len(current_sections)+1}",
                    page_number=current_page_num,
                    heading=current_heading or f"Page {current_page_num}",
                    clause_number=current_clause,
                    text=sec_text,
                    start_char=0,
                    end_char=len(sec_text)
                )
            )

    if current_page_text_parts:
        page_full_text = "\n\n".join(current_page_text_parts)
        pages.append(Page(page_number=current_page_num, text=page_full_text, sections=current_sections))

    full_text = "\n\n".join(p.text for p in pages)
    return pages, full_text
